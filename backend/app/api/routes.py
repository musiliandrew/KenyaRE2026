from fastapi import APIRouter, HTTPException, File, UploadFile
from fastapi.responses import StreamingResponse
import io
import csv
from typing import List, Dict, Any, Optional
from app.models.schemas import (
    PortfolioSummary, ReturnPeriodMetric, EPCurveResponse,
    HazardLookupRequest, HazardLookupResponse,
    DamageCalculationRequest, DamageCalculationResponse,
    RunModelRequest, RunModelResponse,
    NLPParseRequest, NLPParseResponse,
    BriefingRequest, BriefingResponse,
    ExposureUpload, ExposureAsset,
    XoLTreatyRequest, XoLTreatyResponse,
    FacultativeQuoteRequest, FacultativeQuoteResponse,
    AIChatRequest
)
from app.services.financial_engine import financial_engine
from app.services.cat_engine import (
    calculate_jrc_damage_ratio, calculate_asset_loss,
    calculate_portfolio_loss, compute_ep_curve, calculate_aal,
    lookup_hazard_score, apply_ai_drainage_adjustments,
    RP_DEPTH_ANCHORS, RP_LABELS
)
from app.services.ai_service import parse_natural_language_portfolio, generate_risk_briefing, stream_ai_chat
from app.services.synthetic_data import get_synthetic_exposure, get_synthetic_hotspots

router = APIRouter()


# ============================================================================
# SYSTEM ENDPOINTS
# ============================================================================

@router.get("/health", tags=["System"])
def health_check() -> Dict[str, str]:
    """Health check endpoint for monitoring."""
    return {"status": "ok", "service": "Kenya Re Catastrophe Modeling API", "version": "1.0.0"}


@router.get("/", tags=["System"])
def root() -> Dict[str, Any]:
    """Root endpoint with API information."""
    return {
        "message": "Kenya Re Catastrophe Risk Intelligence API",
        "version": "1.0.0",
        "docs": "/docs",
        "organization": "Kenya Reinsurance Corporation",
        "challenge": "Team A - Nairobi Urban Surface-Water Flood",
        "endpoints": {
            "hazard": "/api/hazard/lookup",
            "vulnerability": "/api/vulnerability/calculate",
            "portfolio": "/api/portfolio/summary",
            "ep_curve": "/api/curves/ep",
            "run_model": "/api/model/run",
            "ai_parse": "/api/ai/parse-slip",
            "ai_briefing": "/api/ai/briefing"
        }
    }


# ============================================================================
# PILLAR 1: HAZARD LAYER (CLIMADA Compatible & GeoTIFF Driven)
# ============================================================================

from app.services.hazard_engine import hazard_engine

@router.post("/hazard/lookup", response_model=HazardLookupResponse, tags=["Pillar 1: Hazard"])
def lookup_hazard(req: HazardLookupRequest):
    """
    Sub-millisecond raster lookup: samples pluvial susceptibility and calibrated flood depth
    directly from the 5 GeoTIFF rasters using exact affine transform coordinates.
    """
    res = hazard_engine.get_hazard_depth(req.lat, req.lng, req.return_period.value)
    return HazardLookupResponse(
        lat=res["lat"],
        lng=res["lon"],
        return_period=res["return_period"],
        hazard_score=res["hazard_score"],
        depth_m=res["depth_m"],
        tier_label=res["tier_label"]
    )


@router.get("/hazard/hotspots", tags=["Pillar 1: Hazard"])
def get_validated_hotspots(return_period: str = "100y") -> List[Dict[str, Any]]:
    """
    Returns the 24 official county flood hotspots evaluated across the requested return period raster.
    """
    hotspots = get_synthetic_hotspots()
    return hazard_engine.validate_hotspots(hotspots, return_period)


@router.post("/hazard/climada-matrix", tags=["Pillar 1: Hazard"])
def get_climada_matrix(req: List[Dict[str, float]]) -> Dict[str, Any]:
    """
    Constructs a CLIMADA-compatible Hazard structure (intensity matrix, frequency array, centroids)
    for arbitrary coordinates across all 5 calibrated return periods.
    """
    coords = [(item["lat"], item["lng"]) for item in req]
    matrix = hazard_engine.to_climada_hazard_matrix(coords)
    return {
        "climada_compatible": matrix["climada_compatible"],
        "n_events": matrix["n_events"],
        "n_centroids": matrix["n_centroids"],
        "events": matrix["events"],
        "frequency": matrix["frequency"],
        "intensity_dep_m": matrix["intensity_matrix"].tolist()
    }


# ============================================================================
# PILLAR 2: VULNERABILITY LAYER
# ============================================================================

from app.services.vulnerability_engine import vulnerability_engine

@router.post("/vulnerability/calculate", response_model=DamageCalculationResponse, tags=["Pillar 2: Vulnerability"])
def calculate_damage(req: DamageCalculationRequest):
    """
    Calculates damage ratio and financial loss for a single asset given depth and housing class.
    Uses continuous JRC/Huizinga sigmoid depth-damage curves with physical caps.
    """
    res = vulnerability_engine.calculate_loss(req.tiv_kes, req.depth_m, req.housing_class.value)
    return DamageCalculationResponse(
        housing_class=res["housing_class"],
        depth_m=res["depth_m"],
        tiv_kes=res["tiv_kes"],
        damage_ratio=res["damage_ratio"],
        loss_kes=res["loss_kes"],
        damage_cap_pct=res["cap_pct"]
    )


@router.get("/vulnerability/curves", tags=["Pillar 2: Vulnerability"])
def get_vulnerability_curves(max_depth_m: float = 4.0, step_m: float = 0.1) -> Dict[str, Any]:
    """
    Returns discretized evaluation points for all 4 construction classes for UI charting.
    """
    return {
        "curves": vulnerability_engine.get_all_curves(max_depth_m, step_m),
        "parameters": vulnerability_engine.curves
    }


@router.get("/vulnerability/climada-impact", tags=["Pillar 2: Vulnerability"])
def get_climada_impact_functions() -> Dict[str, Any]:
    """
    Exports all 4 depth-damage functions formatted as CLIMADA ImpactFunc objects.
    """
    return vulnerability_engine.to_climada_impact_funcs()


# ============================================================================
# PILLAR 3: EXPOSURE LAYER
# ============================================================================

from app.services.exposure_engine import exposure_engine

@router.get("/portfolio/summary", response_model=PortfolioSummary, tags=["Pillar 3: Exposure"])
def get_portfolio_summary(return_period: str = "100y"):
    """
    Returns executive metrics for Nairobi urban pluvial exposure using verified baseline portfolio.
    """
    stats = exposure_engine.get_summary_statistics()
    ep = financial_engine.calculate_ep_curve()
    point = next((m for m in ep["ep_curve"] if m["return_period"] == return_period), ep["ep_curve"][-1])

    return PortfolioSummary(
        tiv_kes=stats["total_tiv_kes"],
        event_loss_kes=point["gross_loss_kes"],
        aal_kes=ep["aal_gross_kes"],
        asset_count=stats["asset_count"],
        active_rp=return_period,
        hotspot_count=len(get_synthetic_hotspots()),
        pml_100y_kes=ep["pml_100y_kes"],
        loss_ratio=point["loss_ratio"],
    )


@router.get("/exposure/assets", tags=["Pillar 3: Exposure"])
def get_exposure_assets(
    return_period: str = "100y",
    limit: int = 1000,
    offset: int = 0,
    ward: Optional[str] = None,
    housing_class: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Returns geocoded portfolio assets enriched with spatial hazard depth, damage ratio,
    modeled loss and risk level for the requested return period.
    """
    from app.services.vulnerability_engine import vulnerability_engine

    assets = exposure_engine.assets
    if ward:
        assets = [a for a in assets if a.get("ward", "").lower() == ward.lower()]
    if housing_class:
        assets = [a for a in assets if a["housing_class"] == housing_class]

    enriched = exposure_engine.match_spatial_hazard(return_period, assets=assets)
    for a in enriched:
        calc = vulnerability_engine.calculate_loss(a["tiv_kes"], a["depth_m"], a["housing_class"])
        a["damage_ratio"] = calc["damage_ratio"]
        a["loss_kes"] = calc["loss_kes"]
        a["risk_level"] = "high" if calc["damage_ratio"] > 0.35 else "mid" if calc["damage_ratio"] > 0.08 else "low"

    page = enriched[offset: offset + limit]
    return {"total": len(enriched), "offset": offset, "returned": len(page), "return_period": return_period, "assets": page}


def build_asset_dossier_response(
    asset_dict: Dict[str, Any],
    active_rp: str = "100y",
    portfolio_assets: Optional[List[Dict[str, Any]]] = None
) -> Dict[str, Any]:
    loc_id = str(asset_dict.get("loc_id") or asset_dict.get("id") or "NBO-ASSET").strip()
    name = str(asset_dict.get("name") or loc_id)
    ward = str(asset_dict.get("ward") or "Nairobi")
    lat = float(asset_dict.get("lat") or -1.2847)
    lon = float(asset_dict.get("lon") or asset_dict.get("lng") or 36.8247)
    h_class = str(asset_dict.get("housing_class") or "concrete_rcc").lower().strip()
    area_sqm = float(asset_dict.get("floor_area_m2") or asset_dict.get("area_sqm") or 1000.0)
    tiv_kes = float(asset_dict.get("tiv_kes") or 50_000_000.0)

    # 1. Housing class parameters
    from app.services.vulnerability_engine import vulnerability_engine, HousingClass
    hc_enum = vulnerability_engine.get_housing_class_enum(h_class)
    params = vulnerability_engine.curves.get(hc_enum, vulnerability_engine.curves[HousingClass.concrete_rcc])

    # 2. Discrete Vulnerability Curve for this asset's housing class (0 to 4.5m)
    curve_points = vulnerability_engine.generate_curve_points(hc_enum, max_depth_m=4.5, step_m=0.1)

    # 3. Multi-return-period hazard depths, damage ratios, and losses (5y, 10y, 25y, 50y, 100y)
    from app.services.hazard_engine import hazard_engine, HAZARD_TIER_CONFIG
    rp_metrics = []
    override_depth = asset_dict.get("depth_m")

    for rp_key, rp_cfg in HAZARD_TIER_CONFIG.items():
        if override_depth is not None and float(override_depth) >= 0:
            anchor_active = HAZARD_TIER_CONFIG.get(active_rp, {}).get("depth_anchor_m", 2.2)
            ratio = rp_cfg.get("depth_anchor_m", 1.0) / (anchor_active if anchor_active > 0 else 1.0)
            d_m = max(0.0, float(override_depth) * ratio)
        else:
            h_sample = hazard_engine.get_hazard_depth(lat, lon, rp_key)
            d_m = float(h_sample.get("depth_m", 0.0))

        dr = vulnerability_engine.calculate_damage_ratio(d_m, hc_enum)
        loss = tiv_kes * dr
        deductible_pct = 0.10
        deductible_kes = min(loss, tiv_kes * deductible_pct)
        insured_loss = max(0.0, loss - deductible_kes)

        rp_metrics.append({
            "return_period": rp_key,
            "years": rp_cfg["rp_years"],
            "annual_exceedance_prob": rp_cfg["annual_exceedance_prob"],
            "label": rp_cfg["label"],
            "depth_m": round(d_m, 2),
            "damage_ratio": round(dr, 4),
            "damage_ratio_pct": round(dr * 100, 2),
            "loss_kes": round(loss, 2),
            "deductible_kes": round(deductible_kes, 2),
            "insured_loss_kes": round(insured_loss, 2),
        })

    rp_metrics.sort(key=lambda x: x["years"])

    # 4. Integrate Asset AAL (Annual Average Loss)
    aal_gross_kes = 0.0
    for i in range(len(rp_metrics) - 1):
        p1 = rp_metrics[i]["annual_exceedance_prob"]
        p2 = rp_metrics[i + 1]["annual_exceedance_prob"]
        l1 = rp_metrics[i]["loss_kes"]
        l2 = rp_metrics[i + 1]["loss_kes"]
        delta_p = abs(p1 - p2)
        aal_gross_kes += ((l1 + l2) / 2.0) * delta_p
    if rp_metrics:
        aal_gross_kes += rp_metrics[-1]["loss_kes"] * rp_metrics[-1]["annual_exceedance_prob"] * 0.5

    pure_risk_rate_pct = (aal_gross_kes / tiv_kes * 100.0) if tiv_kes > 0 else 0.0

    active_metric = next((m for m in rp_metrics if m["return_period"] == active_rp), rp_metrics[-1])
    active_depth = active_metric["depth_m"]
    active_damage_ratio = active_metric["damage_ratio"]
    active_loss = active_metric["loss_kes"]
    risk_level = "high" if active_damage_ratio > 0.35 else "mid" if active_damage_ratio > 0.08 else "low"

    # 5. Exposure Portfolio Context & Neighbors
    from app.services.exposure_engine import exposure_engine
    pool = portfolio_assets if portfolio_assets is not None else exposure_engine.assets
    total_portfolio_tiv = sum(a.get("tiv_kes", 0) for a in pool) or tiv_kes
    total_portfolio_loss = sum(a.get("loss_kes", 0) for a in pool) or active_loss

    tiv_share_pct = (tiv_kes / total_portfolio_tiv * 100.0) if total_portfolio_tiv > 0 else 0.0
    loss_share_pct = (active_loss / total_portfolio_loss * 100.0) if total_portfolio_loss > 0 else 0.0

    # Neighbors in same ward
    ward_assets = [
        a for a in pool
        if (a.get("ward") or "").lower() == ward.lower()
        and str(a.get("loc_id") or a.get("id") or "").lower() != loc_id.lower()
    ]
    ward_assets.sort(key=lambda a: a.get("loss_kes", 0) or a.get("tiv_kes", 0), reverse=True)
    top_neighbors = ward_assets[:6]

    return {
        "asset": {
            "loc_id": loc_id,
            "name": name,
            "ward": ward,
            "lat": lat,
            "lon": lon,
            "housing_class": hc_enum.value,
            "housing_class_label": params["label"],
            "floor_area_m2": area_sqm,
            "tiv_kes": tiv_kes,
            "depth_m": active_depth,
            "damage_ratio": active_damage_ratio,
            "damage_ratio_pct": round(active_damage_ratio * 100, 2),
            "loss_kes": active_loss,
            "risk_level": risk_level,
            "active_rp": active_rp,
            "source_file": asset_dict.get("source_file"),
        },
        "financial_summary": {
            "tiv_kes": tiv_kes,
            "active_loss_kes": active_loss,
            "damage_ratio_pct": round(active_damage_ratio * 100, 2),
            "aal_gross_kes": round(aal_gross_kes, 2),
            "pure_risk_rate_pct": round(pure_risk_rate_pct, 4),
            "recommended_premium_kes": round(aal_gross_kes * 1.32, 2),
            "recommended_deductible_pct": 10.0,
            "tiv_share_pct": round(tiv_share_pct, 4),
            "loss_share_pct": round(loss_share_pct, 4),
        },
        "exceedance_probability_curve": rp_metrics,
        "vulnerability_curve": {
            "housing_class": hc_enum.value,
            "label": params["label"],
            "parameters": {
                "cap": params["cap"],
                "cap_pct": round(params["cap"] * 100, 1),
                "k": params["k"],
                "midpoint": params["midpoint"],
                "threshold_m": params["threshold_m"],
                "jrc_reference": params["jrc_reference"],
                "typical_costs_sqm": params["typical_costs_sqm"],
                "description": params["description"],
            },
            "points": curve_points,
            "active_operating_point": {
                "depth_m": active_depth,
                "damage_ratio": active_damage_ratio,
                "damage_ratio_pct": round(active_damage_ratio * 100, 2),
            }
        },
        "top_exposed_locations": [
            {
                "loc_id": a.get("loc_id") or a.get("id"),
                "name": a.get("name") or a.get("loc_id"),
                "ward": a.get("ward"),
                "housing_class": a.get("housing_class"),
                "tiv_kes": a.get("tiv_kes"),
                "depth_m": a.get("depth_m", 0.0),
                "loss_kes": a.get("loss_kes", 0.0),
                "damage_ratio": a.get("damage_ratio", 0.0),
            }
            for a in top_neighbors
        ]
    }


@router.get("/exposure/asset/{loc_id}", tags=["Pillar 3: Exposure"])
def get_single_asset_dossier(loc_id: str, return_period: str = "100y") -> Dict[str, Any]:
    """
    Returns complete single-asset catastrophe dossier including JRC vulnerability curve,
    calibrated parameters, loss & exceedance probability curve, damage ratio, and top exposed neighbors.
    Matches IDs flexibly (e.g. nbo0023, NBO-0023, UH-001).
    """
    clean_id = loc_id.strip().upper().replace(" ", "").replace("_", "-")
    found = None
    for a in exposure_engine.assets:
        aid = str(a.get("loc_id") or a.get("id") or "").upper().replace(" ", "").replace("_", "-")
        if aid == clean_id or aid.replace("-", "") == clean_id.replace("-", ""):
            found = a
            break

    if not found:
        # Fallback search by partial containment
        for a in exposure_engine.assets:
            aid = str(a.get("loc_id") or a.get("id") or "").upper()
            if clean_id in aid or aid in clean_id:
                found = a
                break

    if not found:
        raise HTTPException(
            status_code=404,
            detail=f"Asset ID '{loc_id}' not found in baseline portfolio. Please provide asset data via POST /api/exposure/asset/dossier"
        )

    return build_asset_dossier_response(found, active_rp=return_period)


@router.post("/exposure/asset/dossier", tags=["Pillar 3: Exposure"])
def calculate_asset_dossier(payload: Dict[str, Any]) -> Dict[str, Any]:
    """
    Calculates complete single-asset catastrophe dossier for any arbitrary or uploaded asset.
    """
    asset = payload.get("asset") or payload
    active_rp = str(payload.get("return_period") or asset.get("active_rp") or "100y")
    portfolio_assets = payload.get("portfolio_assets")
    return build_asset_dossier_response(asset, active_rp=active_rp, portfolio_assets=portfolio_assets)


@router.get("/hazard/grid", tags=["Pillar 1: Hazard"])
def get_hazard_grid(return_period: str = "100y", step: int = 12, min_depth_m: float = 0.05) -> Dict[str, Any]:
    """
    Downsampled wet-cell grid from the GeoTIFF raster for the deck.gl / MapLibre flood layer.
    """
    return hazard_engine.get_hazard_grid(return_period, step=max(2, min(step, 64)), min_depth_m=min_depth_m)


@router.get("/exposure/stats", tags=["Pillar 3: Exposure"])
def get_exposure_statistics() -> Dict[str, Any]:
    """
    Returns capital allocation by housing class, TIV share, and administrative ward distribution.
    """
    return exposure_engine.get_summary_statistics()


@router.get("/exposure/climada-entity", tags=["Pillar 3: Exposure"])
def get_climada_exposure_entity() -> Dict[str, Any]:
    """
    Exports full portfolio formatted as CLIMADA Exposures / Entity data dictionary.
    """
    return exposure_engine.to_climada_exposure()


@router.post("/portfolio/upload", tags=["Pillar 3: Exposure"])
def upload_exposure(upload: ExposureUpload):
    """
    Uploads exposure data (Oasis OED format), validates bounding boxes, and returns normalized portfolio.
    """
    raw_dicts = [a.model_dump() for a in upload.assets]
    valid, errors = exposure_engine.validate_exposure_records(raw_dicts)
    return {
        "message": f"Successfully ingested {len(valid)} valid assets",
        "valid_count": len(valid),
        "rejected_count": len(errors),
        "errors": errors,
        "total_tiv_kes": sum(a["tiv_kes"] for a in valid),
        "source": upload.source
    }


@router.get("/portfolio/assets", tags=["Pillar 3: Exposure"])
def get_portfolio_assets(limit: int = 100):
    """
    Returns the baseline exposure portfolio (600 buildings) from the Exposure Engine.
    """
    assets = exposure_engine.assets
    return {
        "total": len(assets),
        "returned": min(limit, len(assets)),
        "assets": assets[:limit]
    }


@router.get("/portfolio/hotspots", tags=["Pillar 3: Exposure"])
def get_hotspots(return_period: str = "100y"):
    """
    Returns county flood hotspots evaluated against the hazard raster (depth and susceptibility).
    """
    hotspots = hazard_engine.validate_hotspots(get_synthetic_hotspots(), return_period)
    return {
        "count": len(hotspots),
        "hotspots": hotspots
    }


# ============================================================================
# PILLAR 4: FINANCIAL ENGINE
# ============================================================================

@router.get("/curves/ep", response_model=EPCurveResponse, tags=["Pillar 4: Financial Engine"])
def get_ep_curve(ai_enabled: bool = True):
    """
    Returns Exceedance Probability (EP) curve data and trapezoidal AAL across all 5 calibrated return periods.
    Powered by the Module 4 Financial Engine with GeoTIFF raster event sampling.
    """
    base = financial_engine.calculate_ep_curve(apply_ai_drainage=False)
    ai = financial_engine.calculate_ep_curve(apply_ai_drainage=True) if ai_enabled else None
    ai_by_rp = {m["return_period"]: m for m in ai["ep_curve"]} if ai else {}

    metrics = []
    for m in base["ep_curve"]:
        a = ai_by_rp.get(m["return_period"])
        metrics.append(ReturnPeriodMetric(
            return_period=m["return_period"],
            years=m["years"],
            annual_prob=m["annual_prob"],
            portfolio_loss_kes=m["gross_loss_kes"],
            loss_ratio=m["loss_ratio"],
            pml_90=m["pml_90"],
            ai_adjusted_loss=a["gross_loss_kes"] if a else None,
            ai_delta_kes=round(a["gross_loss_kes"] - m["gross_loss_kes"], 2) if a else None,
        ))
    return EPCurveResponse(
        metrics=metrics,
        total_tiv_kes=base["total_tiv_kes"],
        aal_kes=ai["aal_gross_kes"] if ai else base["aal_gross_kes"],
        baseline_aal_kes=base["aal_gross_kes"],
        ai_enabled=ai_enabled
    )


@router.post("/reinsurance/xol", response_model=XoLTreatyResponse, tags=["Pillar 4: Financial Engine"])
def price_xol_treaty(req: XoLTreatyRequest):
    """
    Prices an Excess of Loss (XOL) Reinsurance Treaty layer with pure burn rate,
    capital margin, and return-period exhaustion tracking.
    """
    res = financial_engine.price_reinsurance_xol(
        attachment_kes=req.attachment_kes,
        limit_kes=req.limit_kes,
        share_pct=req.share_pct,
        apply_ai_drainage=req.apply_ai_drainage
    )
    return XoLTreatyResponse(**res)


@router.post("/quotes/facultative", response_model=FacultativeQuoteResponse, tags=["Pillar 4: Financial Engine"])
def quote_facultative_slip(req: FacultativeQuoteRequest):
    """
    Calculates technical pure risk premium, deductible retention, and underwriting rates
    for an individual policy slip (e.g. Landmark Plaza Upper Hill from testData.md).
    """
    res = financial_engine.quote_single_slip_facultative(
        tiv_kes=req.tiv_kes,
        lat=req.lat,
        lon=req.lng,
        housing_class=req.housing_class.value,
        deductible_pct=req.deductible_pct
    )
    return FacultativeQuoteResponse(
        tiv_kes=res["tiv_kes"],
        lat=res["lat"],
        lon=res["lon"],
        housing_class=res["housing_class"],
        deductible_pct=res["deductible_pct"],
        deductible_kes=res["deductible_kes"],
        asset_aal_gross_kes=res["asset_aal_gross_kes"],
        asset_aal_insured_kes=res["asset_aal_insured_kes"],
        pure_rate_pct=res["pure_rate_pct"],
        recommended_technical_rate_pct=res["recommended_technical_rate_pct"],
        recommended_annual_premium_kes=res["recommended_annual_premium_kes"],
        depth_100y_m=res["100y_extreme_depth_m"],
        insured_loss_100y_kes=res["100y_extreme_loss_kes"]
    )


@router.post("/model/run", response_model=RunModelResponse, tags=["Pillar 4: Financial Engine"])
def run_model(req: RunModelRequest):
    """
    Runs the full catastrophe pipeline (hazard rasters -> JRC vulnerability -> exposure -> financial engine)
    for a scenario. Uses the uploaded exposure if supplied, otherwise the baseline 600-building portfolio.
    """
    custom = None
    if req.exposure:
        custom = [
            {
                "loc_id": a.id, "name": a.name, "lat": a.lat, "lon": a.lng,
                "housing_class": a.housing_class.value, "floor_area_m2": a.area_sqm,
                "tiv_kes": a.tiv_kes, "ward": a.ward or "Uploaded",
            }
            for a in req.exposure
        ]

    base = financial_engine.calculate_ep_curve(assets=custom, apply_ai_drainage=False)
    active = financial_engine.calculate_ep_curve(assets=custom, apply_ai_drainage=req.apply_ai) if req.apply_ai else base

    rp = req.scenario.value
    elt = next(e for e in active["event_loss_table"] if e["return_period"] == rp)
    base_elt = next(e for e in base["event_loss_table"] if e["return_period"] == rp)
    base_curve = {m["return_period"]: m for m in base["ep_curve"]}

    metrics = []
    for m in active["ep_curve"]:
        b = base_curve[m["return_period"]]
        metrics.append(ReturnPeriodMetric(
            return_period=m["return_period"], years=m["years"], annual_prob=m["annual_prob"],
            portfolio_loss_kes=b["gross_loss_kes"], loss_ratio=b["loss_ratio"], pml_90=b["pml_90"],
            ai_adjusted_loss=m["gross_loss_kes"] if req.apply_ai else None,
            ai_delta_kes=round(m["gross_loss_kes"] - b["gross_loss_kes"], 2) if req.apply_ai else None,
        ))

    return RunModelResponse(
        scenario=rp,
        total_tiv_kes=active["total_tiv_kes"],
        portfolio_loss_kes=elt["gross_portfolio_loss_kes"],
        loss_ratio=elt["gross_loss_ratio"],
        aal_kes=active["aal_gross_kes"],
        asset_count=len(custom) if custom is not None else len(exposure_engine.assets),
        ai_enabled=req.apply_ai,
        ai_delta_kes=round(elt["gross_portfolio_loss_kes"] - base_elt["gross_portfolio_loss_kes"], 2) if req.apply_ai else None,
        aal_ai_delta_kes=round(active["aal_gross_kes"] - base["aal_gross_kes"], 2) if req.apply_ai else None,
        loss_by_class=elt["loss_by_class"],
        loss_by_ward=elt["top_loss_wards"],
        top_losses=elt["top_properties"],
        ep_curve=metrics,
        source="uploaded_exposure" if custom is not None else "baseline_portfolio",
    )



# ============================================================================
# AI INTELLIGENCE LAYER
# ============================================================================

@router.post("/ai/parse-slip", response_model=NLPParseResponse, tags=["AI Intelligence Layer"])
def parse_slip(req: NLPParseRequest):
    """
    Natural Language Exposure Ingestion: parses free-text underwriting notes into structured exposure.
    
    Example: "5 iron sheet warehouses in Mathare along Juja road, 350 sqm each, KES 18M total"
    """
    result = parse_natural_language_portfolio(req.text)
    
    return NLPParseResponse(
        extracted_structures=result["extracted_structures"],
        housing_class=result["housing_class"],
        location=result["location"],
        total_area_sqm=result["total_area_sqm"],
        estimated_tiv_kes=result["estimated_tiv_kes"],
        hazard_tier=result["hazard_tier"],
        estimated_depth_m=result["estimated_depth_m"],
        damage_ratio=result["damage_ratio"],
        recommended_premium_kes=result["recommended_premium_kes"],
        technical_rate_pct=result["technical_rate_pct"],
        summary=result["summary"],
        parsed_assets=[ExposureAsset(**a) for a in result["parsed_assets"]]
    )

@router.post("/ai/briefing", response_model=BriefingResponse, tags=["AI Intelligence Layer"])
def generate_briefing(req: BriefingRequest):
    """
    Generates a natural-language risk briefing for underwriters.
    Summarizes model results in plain English.
    """
    result = generate_risk_briefing(
        scenario=req.scenario.value,
        portfolio_loss_kes=req.portfolio_loss_kes,
        loss_ratio=req.loss_ratio,
        key_hotspots=req.key_hotspots,
        ai_enabled=req.ai_enabled
    )
    
    return BriefingResponse(**result)


@router.post("/ai/chat/stream", tags=["AI Intelligence Layer"])
def chat_stream(req: AIChatRequest):
    """
    Streams interactive AI Copilot responses token-by-token using Groq (openai/gpt-oss-120b).
    """
    def event_generator():
        for token in stream_ai_chat(req.message, req.history):
            yield token

    return StreamingResponse(event_generator(), media_type="text/plain")


def extract_text_from_file(filename: str, content: bytes) -> str:
    """
    Extracts plain text from unstructured files (.docx, .pdf, .txt, .md, .csv).
    """
    ext = filename.lower().split(".")[-1] if "." in filename else ""

    if ext == "pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages).strip()
        except Exception as e:
            return content.decode("utf-8", errors="ignore")

    elif ext in ("docx", "doc"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts).strip()
        except Exception:
            try:
                import zipfile
                import xml.etree.ElementTree as ET
                with zipfile.ZipFile(io.BytesIO(content)) as z:
                    xml_content = z.read("word/document.xml")
                    tree = ET.fromstring(xml_content)
                    return " ".join(node.text for node in tree.iter() if node.text).strip()
            except Exception:
                return content.decode("utf-8", errors="ignore")

    else:
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1", errors="ignore")


@router.post("/ai/ingest-file", tags=["AI Intelligence Layer"])
async def ingest_unstructured_file(file: UploadFile = File(...)):
    """
    Ingests unstructured data files (.docx, .pdf, .txt, .csv, .json, etc.)
    Extracts text and runs AI NLP parsing to extract geocoded CAT portfolio assets.
    """
    content = await file.read()
    filename = file.filename or "uploaded_file"
    ext = filename.lower().split(".")[-1] if "." in filename else ""

    # 1. If it's a structured CSV with standard tabular exposure columns, try tabular parse first
    if ext == "csv":
        try:
            text = content.decode("utf-8")
            lines = [l.strip() for l in text.split("\n") if l.strip()]
            if len(lines) > 1:
                headers = [h.strip().lower() for h in lines[0].split(",")]
                has_lat = any("lat" in h for h in headers)
                has_lon = any("lon" in h or "lng" in h for h in headers)
                has_tiv = any("tiv" in h or "value" in h for h in headers)
                if has_lat and has_lon and has_tiv:
                    reader = csv.DictReader(lines)
                    items = []
                    for i, row in enumerate(reader, 1):
                        row_lower = {k.lower().strip(): v.strip() for k, v in row.items() if k}
                        lat = float(row_lower.get("lat", -1.2921) or -1.2921)
                        lon = float(row_lower.get("lon", row_lower.get("lng", 36.8219)) or 36.8219)
                        tiv = float(row_lower.get("tiv_kes", row_lower.get("tiv", row_lower.get("value", 50000000))) or 50000000)
                        raw_class = row_lower.get("housing_class", row_lower.get("class", "concrete_rcc"))
                        h_class = (
                            "informal_iron_sheet" if "iron" in raw_class
                            else "semi_permanent" if "semi" in raw_class
                            else "permanent_masonry" if "masonry" in raw_class
                            else "concrete_rcc"
                        )
                        name = row_lower.get("name", row_lower.get("building", f"Asset #{i}"))
                        ward = row_lower.get("ward", row_lower.get("locality", "Nairobi"))
                        area = float(row_lower.get("area_sqm", row_lower.get("area", 1000)) or 1000)
                        items.append({
                            "id": f"AST-{i:03d}",
                            "name": name,
                            "lat": lat,
                            "lng": lon,
                            "housing_class": h_class,
                            "area_sqm": area,
                            "tiv_kes": tiv,
                            "ward": ward
                        })
                    if items:
                        return {
                            "file_name": filename,
                            "file_type": ext,
                            "method": "tabular_csv",
                            "extracted_text_preview": f"Tabular exposure CSV with {len(items)} assets",
                            "extracted_structures": len(items),
                            "housing_class": items[0]["housing_class"],
                            "location": items[0]["ward"],
                            "total_area_sqm": sum(a["area_sqm"] for a in items),
                            "estimated_tiv_kes": sum(a["tiv_kes"] for a in items),
                            "parsed_assets": items
                        }
        except Exception:
            pass  # Fall through to unstructured AI extraction

    # 2. Unstructured file extraction (Word DOCX, PDF, Text slip, Markdown, or non-standard file)
    extracted_text = extract_text_from_file(filename, content)
    if not extracted_text or len(extracted_text.strip()) < 5:
        raise HTTPException(status_code=400, detail="Could not extract readable text from uploaded file.")

    result = parse_natural_language_portfolio(extracted_text)

    # Format parsed assets
    assets = []
    if result.get("parsed_assets") and len(result["parsed_assets"]) > 0:
        for i, a in enumerate(result["parsed_assets"], 1):
            assets.append({
                "id": a.get("loc_id") or f"UPL-{i:03d}",
                "name": a.get("name") or f"Extracted Asset #{i}",
                "lat": float(a.get("lat") or -1.2921),
                "lng": float(a.get("lon") or a.get("lng") or 36.8219),
                "housing_class": a.get("housing_class", "concrete_rcc"),
                "area_sqm": float(a.get("floor_area_m2") or 1000),
                "tiv_kes": float(a.get("tiv_kes") or 50000000),
                "ward": a.get("ward") or result.get("location") or "Nairobi",
            })
    else:
        count = max(1, result.get("extracted_structures", 5))
        total_tiv = float(result.get("estimated_tiv_kes", 100000000))
        h_class = result.get("housing_class", "concrete_rcc")
        loc = result.get("location", "Nairobi")
        total_area = float(result.get("total_area_sqm", 5000))
        single_tiv = total_tiv / count
        single_area = total_area / count
        for i in range(count):
            assets.append({
                "id": f"UPL-{i+1:03d}",
                "name": f"{loc} Unit #{i+1}",
                "lat": -1.2995 + (i - count / 2) * 0.002,
                "lng": 36.8152 + (0.001 if i % 2 == 0 else -0.001),
                "housing_class": h_class,
                "area_sqm": single_area,
                "tiv_kes": single_tiv,
                "ward": loc,
            })

    return {
        "file_name": filename,
        "file_type": ext,
        "method": "unstructured_nlp",
        "extracted_text_preview": extracted_text[:300] + ("..." if len(extracted_text) > 300 else ""),
        "extracted_structures": len(assets),
        "housing_class": result.get("housing_class", "concrete_rcc"),
        "location": result.get("location", "Nairobi"),
        "total_area_sqm": result.get("total_area_sqm", sum(a["area_sqm"] for a in assets)),
        "estimated_tiv_kes": result.get("estimated_tiv_kes", sum(a["tiv_kes"] for a in assets)),
        "summary": result.get("summary", ""),
        "parsed_assets": assets
    }


